from docx import Document
from services.ingestion.image import extract_image
import tempfile


def extract_docx(file_path: str) -> list[str]:
    """
    Extract text, tables, and embedded images from DOCX.
    Returns a list of text strings.
    """

    document = Document(file_path)
    extracted_text = []

    # Extract paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted_text.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                extracted_text.append(" | ".join(row_text))

    # Extract embedded images
    for rel in document.part.rels.values():

        if "image" in rel.target_ref:

            image_bytes = rel.target_part.blob

            with tempfile.NamedTemporaryFile(suffix=".jpg") as temp_image:

                temp_image.write(image_bytes)
                temp_image.flush()

                image_text = extract_image(temp_image.name)

                extracted_text.extend(image_text)

    return extracted_text