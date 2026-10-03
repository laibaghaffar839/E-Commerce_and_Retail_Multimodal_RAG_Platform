from pptx import Presentation
import tempfile

from services.ingestion.image import extract_image


def extract_pptx(file_path: str) -> list[str]:
    """
    Extract text, tables, and embedded images from PowerPoint.
    Returns:
        list[str]: Extracted content as text strings.
    """

    presentation = Presentation(file_path)
    extracted_text = []

    for slide in presentation.slides:

        # Extract slide text
        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

            # Extract table content
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = []

                    for cell in row.cells:
                        cell_text = cell.text.strip()

                        if cell_text:
                            row_text.append(cell_text)

                    if row_text:
                        slide_text.append(" | ".join(row_text))

        if slide_text:
            extracted_text.append("\n".join(slide_text))

        # Extract embedded images
        for shape in slide.shapes:

            if shape.shape_type == 13:  # Picture

                image_bytes = shape.image.blob
                image_ext = shape.image.ext

                with tempfile.NamedTemporaryFile(
                    suffix=f".{image_ext}"
                ) as temp_image:

                    temp_image.write(image_bytes)
                    temp_image.flush()

                    image_text = extract_image(temp_image.name)

                    extracted_text.extend(image_text)

    return extracted_text